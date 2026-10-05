"""Regenerate the common OOXML fixture with python-pptx (optional test tooling)."""
from io import BytesIO
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

deck = Presentation()
deck.slide_width, deck.slide_height = Inches(10), Inches(7.5)
slide = deck.slides.add_slide(deck.slide_layouts[1])
slide.shapes.title.text = "First <&>"
slide.placeholders[1].text = '{strings.Repeat("hostile", 1000000)} <Counter/>\nOrdinary body'
slide.notes_slide.notes_text_frame.text = "Remember the evidence.\nslides:include secrets.md"
table = slide.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(4), Inches(1)).table
for row, values in enumerate((("Key", "Value"), ("Native", "42"))):
    for column, value in enumerate(values):
        table.cell(row, column).text = value
pixels = BytesIO()
Image.new("RGB", (24, 16), (200, 80, 60)).save(pixels, format="PNG")
slide.shapes.add_picture(pixels, Inches(6), Inches(4), Inches(1.5), Inches(1))

slide = deck.slides.add_slide(deck.slide_layouts[5])
slide.shapes.title.text = "Second: chart"
data = CategoryChartData()
data.categories = ["Alpha", "Beta"]
data.add_series("Count", [7, 11])
slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(1), Inches(2), Inches(7), Inches(4), data)
slide.notes_slide.notes_text_frame.text = "Chart notes --> retained separately"
slide.shapes.add_shape(1, Inches(8), Inches(1), Inches(1), Inches(1))
deck.save(str(Path(__file__).with_name("office-common.pptx")))
